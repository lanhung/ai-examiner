from __future__ import annotations

import hashlib
import json
import re
import subprocess
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from threading import BoundedSemaphore
from typing import Any

from fastapi import UploadFile
from PIL import Image, ImageDraw, ImageFont


@dataclass
class EvidenceDraft:
    kind: str
    page_number: int
    sequence: int
    label: str = ""
    text: str = ""
    bbox: list[float] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)
    storage_path: str | None = None
    mime_type: str | None = None
    sha256: str = ""


@dataclass
class ParsedDocument:
    text: str
    page_map: list[dict]
    warnings: list[str]
    sha256: str
    evidence: list[EvidenceDraft] = field(default_factory=list)
    document_kind: str = "text"


ALLOWED_SUFFIXES = {".pdf", ".txt", ".md", ".markdown", ".pptx", ".docx"}
_PDF_SLOTS = BoundedSemaphore(2)


def parse_uploaded_document(
    path: Path, raw: bytes, max_chars: int, *, evidence_output_dir: Path,
    max_pdf_pages: int = 500, timeout_seconds: int = 180,
) -> ParsedDocument:
    if path.suffix.lower() not in {".pdf", ".docx", ".pptx"}:
        return parse_document(path, raw, max_chars, evidence_output_dir=evidence_output_dir)
    # A subprocess also works inside Celery's daemonic prefork workers, unlike
    # multiprocessing pools. Keep PyMuPDF isolated from ASGI/worker threads.
    evidence_output_dir.mkdir(parents=True, exist_ok=True)
    result_path = evidence_output_dir / "parsed.json"
    if not _PDF_SLOTS.acquire(timeout=min(timeout_seconds, 10)):
        raise ValueError("材料解析繁忙，请稍后重试；文件尚未解析")
    try:
        try:
            result = subprocess.run(
                [sys.executable, "-m", "ai_examiner.document_parser", str(path.resolve()),
                 str(evidence_output_dir.resolve()), str(max_chars), str(max_pdf_pages)],
                capture_output=True, timeout=timeout_seconds, check=False,
            )
            if not result_path.is_file():
                raise ValueError("材料解析进程失败，请检查文件完整性后重试")
            data = json.loads(result_path.read_text(encoding="utf-8"))
            if result.returncode or "error" in data:
                raise ValueError(data.get("error") or "材料解析失败")
            data["evidence"] = [EvidenceDraft(**item) for item in data["evidence"]]
            return ParsedDocument(**data)
        except subprocess.TimeoutExpired as exc:
            raise ValueError(f"材料解析超过 {timeout_seconds} 秒，请按章节拆分后上传") from exc
        finally:
            result_path.unlink(missing_ok=True)
    finally:
        _PDF_SLOTS.release()


def validate_office_archive(path: Path) -> None:
    """Bound decompression before handing an OOXML archive to a document library."""
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 10000 or sum(item.file_size for item in members) > 200 * 1024 * 1024:
                raise ValueError("Office 文件解压后过大，请拆分材料或压缩内嵌图片")
            if any(item.file_size > 100 * 1024 * 1024 or item.flag_bits & 1 for item in members):
                raise ValueError("Office 文件包含过大的资源或加密内容")
    except zipfile.BadZipFile as exc:
        raise ValueError("Office 文件损坏或不是有效的 DOCX/PPTX 文件") from exc

CJK_FONT_CANDIDATES = (
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJKsc-Regular.otf",
    "C:/Windows/Fonts/msyh.ttc",
    "C:/Windows/Fonts/simhei.ttf",
    "C:/Windows/Fonts/simsun.ttc",
    "/System/Library/Fonts/PingFang.ttc",
    "/System/Library/Fonts/Hiragino Sans GB.ttc",
)
FALLBACK_FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
)


def safe_filename(filename: str) -> str:
    name = Path(filename).name
    name = re.sub(r"[^A-Za-z0-9._\-\u4e00-\u9fff]", "_", name)
    suffix = Path(name).suffix
    return name[:180 - len(suffix)] + suffix if len(name) > 180 else name


def save_upload(
    upload: UploadFile, destination: Path, max_bytes: int, *, original_filename: str | None = None
) -> tuple[Path, bytes]:
    # Called from the upload route's worker thread, never the ASGI event loop.
    raw = upload.file.read(max_bytes + 1)
    if len(raw) > max_bytes:
        raise ValueError(f"文件超过 {max_bytes // (1024 * 1024)} MB 限制")
    filename = original_filename or upload.filename or "upload"
    suffix = Path(filename).suffix.lower()
    if suffix == ".doc":
        raise ValueError("暂不支持旧版 .doc，请在 Word 中另存为 .docx 后上传")
    if suffix not in ALLOWED_SUFFIXES:
        raise ValueError("支持 PDF、PPTX、DOCX、TXT 和 Markdown")
    destination.mkdir(parents=True, exist_ok=True)
    # The caller supplies a unique staging directory. Keep physical names short
    # for Windows paths; the original display name is saved in Document metadata.
    path = destination / f"document{suffix}"
    path.write_bytes(raw)
    return path, raw


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _font(size: int = 18):
    # Logical previews frequently contain Chinese text. Prefer a CJK font before
    # DejaVu so vision models receive readable glyphs instead of tofu boxes.
    for candidate in (*CJK_FONT_CANDIDATES, *FALLBACK_FONT_CANDIDATES):
        try:
            return ImageFont.truetype(candidate, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def _text_preview(lines: list[str], output: Path, title: str, width: int = 1280) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    height = max(720, 130 + min(len(lines), 32) * 30)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((45, 35), title[:100], fill="black", font=_font(28))
    y = 95
    for line in lines[:32]:
        clean = re.sub(r"\s+", " ", line).strip()
        if not clean:
            continue
        while clean:
            part, clean = clean[:95], clean[95:]
            draw.text((50, y), part, fill="black", font=_font(18))
            y += 28
            if y > height - 35:
                break
        if y > height - 35:
            break
    image.save(output, "PNG")


def _parse_pdf(path: Path, evidence_dir: Path, warnings: list[str], max_pages: int = 500) -> tuple[list[str], list[dict], list[EvidenceDraft]]:
    import fitz

    document = fitz.open(path)
    if document.needs_pass:
        document.close()
        raise ValueError("PDF 已加密，请先解除密码保护后上传")
    if len(document) > max_pages:
        document.close()
        raise ValueError(f"PDF 超过 {max_pages} 页限制，请按章节拆分后上传")
    chunks: list[str] = []
    pages: list[dict] = []
    evidence: list[EvidenceDraft] = []
    char_cursor = 0
    for page_index, page in enumerate(document, start=1):
        page_dict = page.get_text("dict")
        blocks = page_dict.get("blocks", [])
        page_text = page.get_text("text").strip()
        page_header = f"\n--- PAGE {page_index} ---\n"
        chunks.append(page_header + page_text)
        page_path = evidence_dir / f"page-{page_index:04d}.png"
        # Bound unusually large page dimensions while preserving ordinary 144-dpi pages.
        scale = min(2.0, (4_000_000 / max(page.rect.width * page.rect.height, 1)) ** 0.5)
        matrix = fitz.Matrix(scale, scale)
        pix = page.get_pixmap(matrix=matrix, alpha=False)
        pix.save(page_path)
        page_width, page_height = float(page.rect.width), float(page.rect.height)
        page_asset = EvidenceDraft(
            kind="page",
            page_number=page_index,
            sequence=0,
            label=f"Page {page_index}",
            text=page_text[:4000],
            bbox=[0.0, 0.0, page_width, page_height],
            metadata={"width": page_width, "height": page_height, "format": "pdf"},
            storage_path=str(page_path),
            mime_type="image/png",
            sha256=_sha(page_path),
        )
        evidence.append(page_asset)
        sequence = 1
        for block in blocks:
            block_type = int(block.get("type", 0))
            bbox = [round(float(v), 2) for v in block.get("bbox", [])]
            if block_type == 0:
                lines = []
                for line in block.get("lines", []):
                    spans = [str(span.get("text", "")) for span in line.get("spans", [])]
                    lines.append("".join(spans))
                text = "\n".join(lines).strip()
                if not text:
                    continue
                lowered = text.lower()
                kind = "text_block"
                if re.search(r"(^|\n)\s*(figure|fig\.|图\s*\d+)", lowered):
                    kind = "figure_caption"
                elif re.search(r"(^|\n)\s*(table|表\s*\d+)", lowered):
                    kind = "table_caption"
                elif any(token in text for token in ("=", "∑", "∫", "\u03b1", "\u03b2")) and len(text) < 800:
                    kind = "formula_context"
                evidence.append(
                    EvidenceDraft(
                        kind=kind,
                        page_number=page_index,
                        sequence=sequence,
                        label=text.splitlines()[0][:120],
                        text=text,
                        bbox=bbox,
                        metadata={"source": "pdf_text_block"},
                    )
                )
                sequence += 1
            elif block_type == 1 and block.get("image"):
                ext = block.get("ext") or "png"
                image_path = evidence_dir / f"page-{page_index:04d}-image-{sequence:03d}.{ext}"
                try:
                    image_path.write_bytes(block["image"])
                    mime = f"image/{'jpeg' if ext in {'jpg', 'jpeg'} else ext}"
                    evidence.append(
                        EvidenceDraft(
                            kind="embedded_image",
                            page_number=page_index,
                            sequence=sequence,
                            label=f"Embedded image on page {page_index}",
                            bbox=bbox,
                            metadata={"source": "pdf_embedded_image"},
                            storage_path=str(image_path),
                            mime_type=mime,
                            sha256=_sha(image_path),
                        )
                    )
                    sequence += 1
                except Exception as exc:  # pragma: no cover - corrupt embedded image
                    warnings.append(f"第 {page_index} 页内嵌图片提取失败：{exc}")
        pages.append(
            {
                "page": page_index,
                "start_char": char_cursor,
                "end_char": char_cursor + len(page_header) + len(page_text),
                "text": page_text[:2000],
                "width": page_width,
                "height": page_height,
                "preview_path": str(page_path),
            }
        )
        char_cursor = pages[-1]["end_char"]
    document.close()
    if sum(len(chunk) for chunk in chunks) < 300:
        warnings.append("PDF 可提取文本很少，可能是扫描件；已生成页面截图，但尚未启用 OCR。")
    return chunks, pages, evidence


def _shape_bbox(shape, slide_width: int, slide_height: int) -> list[float]:
    return [
        round(shape.left / slide_width, 5),
        round(shape.top / slide_height, 5),
        round((shape.left + shape.width) / slide_width, 5),
        round((shape.top + shape.height) / slide_height, 5),
    ]


def _parse_pptx(path: Path, evidence_dir: Path, warnings: list[str]) -> tuple[list[str], list[dict], list[EvidenceDraft]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    chunks: list[str] = []
    pages: list[dict] = []
    evidence: list[EvidenceDraft] = []
    cursor = 0
    slide_w, slide_h = presentation.slide_width, presentation.slide_height
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        slide_evidence: list[EvidenceDraft] = []
        for sequence, shape in enumerate(slide.shapes, start=1):
            bbox = _shape_bbox(shape, slide_w, slide_h)
            if getattr(shape, "has_text_frame", False):
                text = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
                if text:
                    texts.append(text)
                    kind = "slide_title" if sequence == 1 or shape == slide.shapes.title else "text_block"
                    slide_evidence.append(
                        EvidenceDraft(
                            kind=kind,
                            page_number=slide_number,
                            sequence=sequence,
                            label=text.splitlines()[0][:120],
                            text=text,
                            bbox=bbox,
                            metadata={"source": "pptx_shape", "shape_type": str(shape.shape_type)},
                        )
                    )
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image = shape.image
                    image_path = evidence_dir / f"slide-{slide_number:04d}-image-{sequence:03d}.{image.ext}"
                    image_path.write_bytes(image.blob)
                    slide_evidence.append(
                        EvidenceDraft(
                            kind="slide_image",
                            page_number=slide_number,
                            sequence=sequence,
                            label=f"Image on slide {slide_number}",
                            bbox=bbox,
                            metadata={"source": "pptx_picture"},
                            storage_path=str(image_path),
                            mime_type=image.content_type,
                            sha256=_sha(image_path),
                        )
                    )
                except Exception as exc:  # pragma: no cover
                    warnings.append(f"第 {slide_number} 页图片提取失败：{exc}")
            if getattr(shape, "has_table", False):
                rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows]
                table_text = "\n".join(rows)
                texts.append(table_text)
                slide_evidence.append(
                    EvidenceDraft(
                        kind="table",
                        page_number=slide_number,
                        sequence=sequence,
                        label=f"Table on slide {slide_number}",
                        text=table_text,
                        bbox=bbox,
                        metadata={"source": "pptx_table"},
                    )
                )
        notes = ""
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ""
        if notes:
            texts.append(f"Speaker notes:\n{notes}")
            slide_evidence.append(
                EvidenceDraft(
                    kind="speaker_notes",
                    page_number=slide_number,
                    sequence=999,
                    label=f"Speaker notes {slide_number}",
                    text=notes,
                    metadata={"source": "pptx_notes"},
                )
            )
        page_text = "\n".join(texts).strip()
        header = f"\n--- SLIDE {slide_number} ---\n"
        chunks.append(header + page_text)
        preview = evidence_dir / f"slide-{slide_number:04d}.png"
        _text_preview(texts, preview, f"Slide {slide_number}")
        evidence.append(
            EvidenceDraft(
                kind="page",
                page_number=slide_number,
                sequence=0,
                label=f"Slide {slide_number}",
                text=page_text[:4000],
                bbox=[0, 0, 1, 1],
                metadata={"width_emu": slide_w, "height_emu": slide_h, "format": "pptx"},
                storage_path=str(preview),
                mime_type="image/png",
                sha256=_sha(preview),
            )
        )
        evidence.extend(slide_evidence)
        pages.append(
            {
                "page": slide_number,
                "label": f"Slide {slide_number}",
                "start_char": cursor,
                "end_char": cursor + len(header) + len(page_text),
                "text": page_text[:2000],
                "preview_path": str(preview),
            }
        )
        cursor = pages[-1]["end_char"]
    return chunks, pages, evidence


def _parse_docx(path: Path, evidence_dir: Path, warnings: list[str]) -> tuple[list[str], list[dict], list[EvidenceDraft]]:
    from docx import Document as DocxDocument
    from docx.table import Table

    document = DocxDocument(path)
    lines: list[str] = []
    evidence: list[EvidenceDraft] = []
    sequence = 1
    table_index = 0
    for paragraph in document.iter_inner_content():
        if isinstance(paragraph, Table):
            table_index += 1
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in paragraph.rows]
            table_text = "\n".join(rows)
            lines.append(f"Table {table_index}:\n{table_text}")
            evidence.append(EvidenceDraft(
                kind="table", page_number=1, sequence=sequence,
                label=f"Table {table_index}", text=table_text,
                metadata={"source": "docx_table"},
            ))
            sequence += 1
            continue
        text = paragraph.text.strip()
        if not text:
            continue
        lines.append(text)
        style_name = getattr(paragraph.style, "name", "") or ""
        evidence.append(
            EvidenceDraft(
                kind="heading" if style_name.lower().startswith("heading") else "text_block",
                page_number=1,
                sequence=sequence,
                label=text[:120],
                text=text,
                metadata={"source": "docx_paragraph", "style": style_name},
            )
        )
        sequence += 1
    try:
        with zipfile.ZipFile(path) as archive:
            media_names = [name for name in archive.namelist() if name.startswith("word/media/")]
            for media_index, name in enumerate(media_names, start=1):
                suffix = Path(name).suffix or ".bin"
                output = evidence_dir / f"docx-image-{media_index:03d}{suffix}"
                output.write_bytes(archive.read(name))
                evidence.append(
                    EvidenceDraft(
                        kind="document_image",
                        page_number=1,
                        sequence=sequence,
                        label=f"Document image {media_index}",
                        metadata={"source": "docx_media", "archive_name": name},
                        storage_path=str(output),
                        mime_type=f"image/{suffix.lstrip('.').replace('jpg', 'jpeg')}",
                        sha256=_sha(output),
                    )
                )
                sequence += 1
    except Exception as exc:  # pragma: no cover
        warnings.append(f"DOCX 图片提取失败：{exc}")
    text = "\n".join(lines)
    warnings.append("Word 为逻辑预览，不代表 Word 原始分页；表格按正文顺序解析。")
    preview = evidence_dir / "document-page-0001.png"
    _text_preview(lines, preview, path.name)
    evidence.insert(
        0,
        EvidenceDraft(
            kind="page",
            page_number=1,
            sequence=0,
            label="Document preview",
            text=text[:4000],
            bbox=[0, 0, 1, 1],
            metadata={"format": "docx", "note": "logical preview; DOCX pagination is renderer-dependent"},
            storage_path=str(preview),
            mime_type="image/png",
            sha256=_sha(preview),
        ),
    )
    return [text], [{"page": 1, "start_char": 0, "end_char": len(text), "text": text[:2000], "preview_path": str(preview)}], evidence


def _parse_plain(path: Path, raw: bytes, evidence_dir: Path, warnings: list[str]) -> tuple[list[str], list[dict], list[EvidenceDraft]]:
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("utf-8", errors="replace")
        warnings.append("文件包含无法按 UTF-8 解码的字符，已进行替换。")
    preview = evidence_dir / "text-page-0001.png"
    _text_preview(text.splitlines(), preview, path.name)
    evidence = [
        EvidenceDraft(
            kind="page",
            page_number=1,
            sequence=0,
            label="Text document",
            text=text[:4000],
            bbox=[0, 0, 1, 1],
            metadata={"format": path.suffix.lower().lstrip(".")},
            storage_path=str(preview),
            mime_type="image/png",
            sha256=_sha(preview),
        ),
        EvidenceDraft(
            kind="text_block",
            page_number=1,
            sequence=1,
            label=text.splitlines()[0][:120] if text.splitlines() else path.name,
            text=text,
            metadata={"source": "plain_text"},
        ),
    ]
    return [text], [{"page": 1, "start_char": 0, "end_char": len(text), "text": text[:2000], "preview_path": str(preview)}], evidence


def parse_document(
    path: Path,
    raw: bytes,
    max_chars: int,
    evidence_output_dir: Path | None = None,
    max_pdf_pages: int = 500,
) -> ParsedDocument:
    suffix = path.suffix.lower()
    warnings: list[str] = []
    evidence_dir = evidence_output_dir or path.parent / f"{path.stem}_evidence"
    evidence_dir.mkdir(parents=True, exist_ok=True)

    if suffix in {".docx", ".pptx"}:
        validate_office_archive(path)

    if suffix == ".pdf":
        chunks, pages, evidence = _parse_pdf(path, evidence_dir, warnings, max_pdf_pages)
        kind = "pdf"
    elif suffix == ".pptx":
        chunks, pages, evidence = _parse_pptx(path, evidence_dir, warnings)
        kind = "presentation"
    elif suffix == ".docx":
        chunks, pages, evidence = _parse_docx(path, evidence_dir, warnings)
        kind = "document"
    else:
        chunks, pages, evidence = _parse_plain(path, raw, evidence_dir, warnings)
        kind = "text"

    content = "\n".join(chunks).strip()
    if len(content) > max_chars:
        warnings.append(f"文档超过 {max_chars} 字符，已截断用于模型分析；证据索引仍保留。")
        content = content[:max_chars]
    if not content:
        raise ValueError("未能从文件中提取可用文本")
    return ParsedDocument(
        text=content,
        page_map=pages,
        warnings=warnings,
        sha256=hashlib.sha256(raw).hexdigest(),
        evidence=evidence,
        document_kind=kind,
    )
