from datetime import UTC, datetime, timedelta
from uuid import uuid4

import pytest
from docx import Document as WordDocument
from sqlalchemy import select, update
from test_blueprint_async import _project_and_document
from test_school_exams import ROOT, headers, join, publish
from test_school_exams import classroom as _classroom

from ai_examiner.agents.grounding import GroundingChecker
from ai_examiner.db import SessionLocal
from ai_examiner.models import (
    BackgroundJob,
    Blueprint,
    Document,
    ExamSession,
    SchoolAttempt,
    StoredObject,
)
from ai_examiner.services.documents import parse_document
from ai_examiner.services.jobs import dispatch_job_by_id


@pytest.fixture
def classroom(monkeypatch):
    return _classroom.__wrapped__(monkeypatch)


@pytest.mark.parametrize("quote", ["", "Invented finding", "Accuracy is 99%"])
def test_short_or_empty_evidence_fails(quote):
    result = GroundingChecker().check_blueprint({"questions": [{"id": "Q1", "source_excerpt": quote}]}, "Real findings")
    assert not result["passed"]


def test_evidence_page_must_match_even_when_quote_exists_elsewhere():
    question = {"id": "Q1", "source_excerpt": "独立验证", "source_page": 1}
    result = GroundingChecker().check_blueprint({"questions": [question]},
                                               "--- PAGE 1 ---\n方法介绍\n--- PAGE 2 ---\n独立\n验证",
                                               [{"page": 1}, {"page": 2}])
    assert not result["passed"]
    question["source_page"] = 2
    assert GroundingChecker().check_blueprint({"questions": [question]},
                                              "--- PAGE 1 ---\n方法介绍\n--- PAGE 2 ---\n独立\n验证",
                                              [{"page": 1}, {"page": 2}])["passed"]


def test_publish_rechecks_sources_not_cached_passed_flag(client, classroom):
    with SessionLocal() as db:
        blueprint = db.get(Blueprint, classroom["blueprint"])
        blueprint.data = {**blueprint.data, "grounding_check": {"passed": True}, "questions": [
            {**q, "source_excerpt": "Invented finding"} for q in blueprint.data["questions"]]}
        db.commit()
    response = publish(client, classroom)
    assert response.status_code == 422
    assert "来源校验" in response.text


def test_docx_preserves_body_order_and_warns_about_logical_pages(tmp_path):
    word = WordDocument()
    word.add_paragraph("BEFORE")
    table = word.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "RESULT_TABLE"
    word.add_paragraph("AFTER")
    path = tmp_path / "paper.docx"
    word.save(path)
    parsed = parse_document(path, path.read_bytes(), 10000)
    assert parsed.text.index("BEFORE") < parsed.text.index("RESULT_TABLE") < parsed.text.index("AFTER")
    assert [item.kind for item in parsed.evidence[1:]] == ["text_block", "table", "text_block"]
    assert any("逻辑预览" in warning for warning in parsed.warnings)


def test_async_upload_receipt_precedes_parsing_and_is_idempotent(client, monkeypatch):
    real_dispatch = dispatch_job_by_id
    monkeypatch.setattr("ai_examiner.services.jobs.dispatch_job_by_id", lambda *args: None)
    project = client.post("/api/projects", json={"name": "Upload receipt"}).json()
    key = str(uuid4())
    url = f"/api/projects/{project['id']}/documents/async"
    response = client.post(url, data={"request_id": key}, files={"file": ("paper.txt", b"Synthetic evidence", "text/plain")})
    assert response.status_code == 202, response.text
    job = response.json()
    assert job["status"] == "queued"
    assert client.get(f"/api/projects/{project['id']}/document-task?request_id={key}").json()["job"]["id"] == job["id"]
    duplicate = client.post(url, data={"request_id": key}, files={"file": ("paper.txt", b"Synthetic evidence", "text/plain")})
    assert duplicate.json()["id"] == job["id"]
    conflict = client.post(url, data={"request_id": key}, files={"file": ("paper.txt", b"Changed", "text/plain")})
    assert conflict.status_code == 409
    with SessionLocal() as db:
        assert db.query(Document).count() == 0
    real_dispatch(job["id"])
    result = client.get(f"/api/projects/{project['id']}/document-task?request_id={key}").json()["job"]
    assert result["status"] == "completed", result
    assert result["result"]["document"]["char_count"] == 18
    real_dispatch(job["id"])
    with SessionLocal() as db:
        assert db.query(Document).count() == 1


def make_session(client):
    project, document = _project_and_document(client)
    blueprint = client.post(f"/api/projects/{project['id']}/blueprints?document_id={document['id']}").json()
    response = client.post("/api/sessions", json={"project_id": project["id"], "blueprint_id": blueprint["id"], "profile": "mock:heuristic-v2"})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def test_session_receipt_replays_without_duplicating_turns(client):
    session = make_session(client)
    url = f"/api/sessions/{session}/operations"
    request = {"request_id": str(uuid4()), "operation": "start", "expected_turn_count": 0}
    first = client.post(url, json=request)
    assert first.status_code == 202, first.text
    result = client.get(f"{url}/{request['request_id']}").json()["job"]
    assert result["status"] == "completed", result
    assert client.post(url, json=request).json()["id"] == first.json()["id"]
    assert len(client.get(f"/api/sessions/{session}").json()["turns"]) == 1
    request = {"request_id": str(uuid4()), "operation": "answer", "expected_turn_count": 1,
               "answer": "Independent validation and calibrated evidence, with distribution limitations."}
    first = client.post(url, json=request)
    assert first.status_code == 202, first.text
    result = client.get(f"{url}/{request['request_id']}").json()["job"]
    assert result["status"] == "completed", result
    assert client.post(url, json=request).json()["id"] == first.json()["id"]
    assert client.post(url, json={**request, "answer": "Changed"}).status_code == 409
    assert client.post(url, json={**request, "request_id": str(uuid4())}).status_code == 409
    assert len(client.get(f"/api/sessions/{session}").json()["turns"]) == 3


def test_session_missing_receipt_and_queued_operation_are_distinct(client, monkeypatch):
    monkeypatch.setattr("ai_examiner.services.jobs.dispatch_job_by_id", lambda *args: None)
    session = make_session(client)
    key = str(uuid4())
    url = f"/api/sessions/{session}/operations"
    assert client.get(f"{url}/{key}").json() == {"job": None}
    body = {"request_id": key, "operation": "start", "expected_turn_count": 0}
    queued = client.post(url, json=body).json()
    assert queued["status"] == "queued"
    assert client.post(url, json={**body, "request_id": str(uuid4())}).status_code == 409
    assert client.post(f"/api/sessions/{session}/start").status_code == 409
    assert client.get(f"{url}/{key}").json()["job"]["status"] == "queued"


def test_lost_session_worker_fenced_before_explicit_new_operation(client, monkeypatch):
    monkeypatch.setattr("ai_examiner.services.jobs.dispatch_job_by_id", lambda *args: None)
    session = make_session(client)
    key = str(uuid4())
    url = f"/api/sessions/{session}/operations"
    job = client.post(url, json={"request_id": key, "operation": "start", "expected_turn_count": 0}).json()
    with SessionLocal() as db:
        row = db.get(BackgroundJob, job["id"])
        row.status, row.lease_token = "running", "expired-worker"
        row.lease_expires_at = datetime.now(UTC) - timedelta(seconds=10)
        db.commit()
    assert client.get(f"{url}/{key}").json()["job"]["status"] == "failed"
    with SessionLocal() as db:
        assert db.get(ExamSession, session).state != "OPERATION_QUEUED"
    assert client.post(url, json={"request_id": str(uuid4()), "operation": "start", "expected_turn_count": 0}).status_code == 202


def test_async_student_grading_preserves_ownership_and_exactly_once_score(client, classroom):
    exam = publish(client, classroom, question_count=1).json()
    attempt = join(client, classroom, exam["id"]).json()
    url = f"{ROOT}/attempts/{attempt['id']}/answers/async"
    body = {"position": 0, "request_id": str(uuid4()), "answer": "Hold out test data and check distribution shift."}
    assert client.post(url, json=body, headers=headers(classroom, "stranger")).status_code == 404
    accepted = client.post(url, json=body, headers=headers(classroom, "student"))
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["status"] == "grading"
    result = client.get(f"{ROOT}/attempts/{attempt['id']}", headers=headers(classroom, "student"))
    assert result.json()["status"] == "completed", result.text
    assert client.post(url, json=body, headers=headers(classroom, "student")).json()["status"] == "completed"
    assert len(classroom["calls"]) == 1
    grades = client.get(f"{ROOT}/teacher/exams/{exam['id']}/grades", headers=headers(classroom, "teacher")).json()
    assert grades[0]["score"] == result.json()["score"]


def test_lost_grading_job_becomes_reviewable_not_permanently_grading(client, classroom, monkeypatch):
    monkeypatch.setattr("ai_examiner.services.jobs.dispatch_job_by_id", lambda *args: None)
    exam = publish(client, classroom, question_count=1).json()
    attempt = join(client, classroom, exam["id"]).json()
    body = {"position": 0, "request_id": str(uuid4()), "answer": "Synthetic answer"}
    client.post(f"{ROOT}/attempts/{attempt['id']}/answers/async", json=body, headers=headers(classroom, "student"))
    with SessionLocal() as db:
        job = db.scalar(select(BackgroundJob).where(BackgroundJob.kind == "school_grading"))
        job.status, job.lease_token = "running", "dead-worker"
        job.lease_expires_at = datetime.now(UTC) - timedelta(seconds=10)
        db.commit()
    result = client.get(f"{ROOT}/attempts/{attempt['id']}", headers=headers(classroom, "student")).json()
    assert result["status"] == "needs_review"
    assert "pending_answer" not in result
    with SessionLocal() as db:
        assert db.get(SchoolAttempt, attempt["id"]).pending["answer"] == body["answer"]


def test_grading_admission_retries_without_losing_student_answer(client, classroom, monkeypatch):
    from types import SimpleNamespace

    from ai_examiner import school_api
    from ai_examiner.jobs import execute_job_delivery
    from ai_examiner.services.model_governance import ModelGovernanceError

    monkeypatch.setattr("ai_examiner.services.jobs.dispatch_job_by_id", lambda *args: None)
    exam = publish(client, classroom, question_count=1).json()
    attempt = join(client, classroom, exam["id"]).json()
    body = {"position": 0, "request_id": str(uuid4()), "answer": "Independent validation."}
    client.post(f"{ROOT}/attempts/{attempt['id']}/answers/async", json=body, headers=headers(classroom, "student"))
    original = school_api.governed_provider
    attempts = []

    def capacity(*args, **kwargs):
        attempts.append(1)
        if len(attempts) == 1:
            raise ModelGovernanceError("organization_concurrency_limited", "Busy", status_code=429)
        return original(*args, **kwargs)

    monkeypatch.setattr(school_api, "governed_provider", capacity)
    retry_calls = []
    task = SimpleNamespace(request=SimpleNamespace(id="delivery"),
                           retry=lambda **kwargs: (retry_calls.append(kwargs), RuntimeError("scheduled"))[1])
    with SessionLocal() as db:
        job = db.scalar(select(BackgroundJob).where(BackgroundJob.kind == "school_grading"))
        envelope, job_id = job.envelope_json, job.id
    with pytest.raises(RuntimeError, match="scheduled"):
        execute_job_delivery(task, envelope)
    assert retry_calls[0]["max_retries"] == 29
    with SessionLocal() as db:
        job = db.get(BackgroundJob, job_id)
        assert job.status == "retry_scheduled"
        assert db.get(SchoolAttempt, attempt["id"]).status == "grading"
        assert db.get(SchoolAttempt, attempt["id"]).pending["answer"] == body["answer"]
        job.next_attempt_at = datetime.now(UTC) - timedelta(seconds=1)
        db.commit()
    execute_job_delivery(task, envelope)
    result = client.get(f"{ROOT}/attempts/{attempt['id']}", headers=headers(classroom, "student")).json()
    assert result["status"] == "completed"
    assert len(classroom["calls"]) == 1


def test_grading_unknown_model_outcome_never_auto_replays(client, classroom, monkeypatch):
    from types import SimpleNamespace

    from fastapi import HTTPException

    from ai_examiner.jobs import execute_job_delivery

    monkeypatch.setattr("ai_examiner.services.jobs.dispatch_job_by_id", lambda *args: None)
    exam = publish(client, classroom, question_count=1).json()
    attempt = join(client, classroom, exam["id"]).json()
    body = {"position": 0, "request_id": str(uuid4()), "answer": "Preserve this answer."}
    client.post(f"{ROOT}/attempts/{attempt['id']}/answers/async", json=body, headers=headers(classroom, "student"))
    monkeypatch.setattr("ai_examiner.school_api.AnswerAnalyzer.analyze", lambda *a, **k: (_ for _ in ()).throw(TimeoutError("Unknown remote result")))
    task = SimpleNamespace(request=SimpleNamespace(id="delivery"), retry=lambda **kwargs: pytest.fail("Paid request replayed"))
    with SessionLocal() as db:
        job = db.scalar(select(BackgroundJob).where(BackgroundJob.kind == "school_grading"))
        envelope, job_id = job.envelope_json, job.id
    with pytest.raises(HTTPException):
        execute_job_delivery(task, envelope)
    with SessionLocal() as db:
        assert db.get(BackgroundJob, job_id).status == "failed"
        row = db.get(SchoolAttempt, attempt["id"])
        assert row.status == "needs_review" and row.pending["answer"] == body["answer"]


def test_office_archive_limit_rejected_before_opening_document(tmp_path, monkeypatch):
    from types import SimpleNamespace

    from ai_examiner.services import documents

    class Archive:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def infolist(self):
            return [SimpleNamespace(file_size=201 * 1024 * 1024)]

    monkeypatch.setattr(documents.zipfile, "ZipFile", lambda *a: Archive())
    with pytest.raises(ValueError, match="解压后过大"):
        documents.validate_office_archive(tmp_path / "oversized.docx")


def test_failed_parse_retries_stored_source_without_new_upload(client, monkeypatch):
    from ai_examiner.services import document_ingestion

    original = document_ingestion.parse_uploaded_document
    monkeypatch.setattr(document_ingestion, "parse_uploaded_document", lambda *a, **k: (_ for _ in ()).throw(ValueError("解析繁忙")))
    project = client.post("/api/projects", json={"name": "Parser retry"}).json()
    key = str(uuid4())
    response = client.post(f"/api/projects/{project['id']}/documents/async", data={"request_id": key},
                           files={"file": ("paper.txt", b"Synthetic material", "text/plain")})
    assert response.status_code == 202, response.text
    receipt = f"/api/projects/{project['id']}/document-task?request_id={key}"
    assert client.get(receipt).json()["job"]["status"] == "failed"
    monkeypatch.setattr(document_ingestion, "parse_uploaded_document", original)
    retried = client.post(f"/api/projects/{project['id']}/document-task/{key}/retry")
    assert retried.status_code == 202, retried.text
    result = client.get(receipt).json()["job"]
    assert result["id"] == response.json()["id"] and result["status"] == "completed"
    with SessionLocal() as db:
        assert db.query(Document).count() == 1 and db.query(BackgroundJob).count() == 1


def test_teacher_preview_async_score_never_enters_formal_grade_list(client, classroom):
    exam = publish(client, classroom, question_count=1).json()
    h = headers(classroom, "teacher")
    preview = client.post(f"{ROOT}/teacher/exams/{exam['id']}/preview", headers=h).json()
    response = client.post(f"{ROOT}/teacher/previews/{preview['id']}/answers/async", headers=h,
                           json={"request_id": str(uuid4()), "position": 0, "answer": "Independent validation"})
    assert response.status_code == 200, response.text
    result = client.get(f"{ROOT}/teacher/previews/{preview['id']}", headers=h).json()
    assert result["status"] == "completed" and result["is_preview"]
    assert client.get(f"{ROOT}/teacher/exams/{exam['id']}/grades", headers=h).json() == []


def test_session_model_failure_is_terminal_and_explicit_resubmit_is_possible(client, monkeypatch):
    session = make_session(client)
    url = f"/api/sessions/{session}/operations"
    client.post(url, json={"request_id": str(uuid4()), "operation": "start", "expected_turn_count": 0})
    calls = []

    def timeout(*args, **kwargs):
        calls.append(1)
        raise TimeoutError("Uncertain remote result")

    monkeypatch.setattr("ai_examiner.services.session_operations.ExamOrchestrator.submit_answer", timeout)
    body = {"request_id": str(uuid4()), "operation": "answer", "expected_turn_count": 1, "answer": "Draft"}
    assert client.post(url, json=body).status_code == 202
    receipt = client.get(f"{url}/{body['request_id']}").json()["job"]
    assert receipt["status"] == "failed"
    assert len(calls) == 1
    assert client.post(url, json=body).json()["status"] == "failed"
    assert len(calls) == 1
    assert client.post(url, json={**body, "request_id": str(uuid4())}).status_code == 202
    assert len(calls) == 2


def test_racing_upload_reuses_receipt_and_keeps_only_winning_source(client, monkeypatch):
    from ai_examiner.config import get_settings
    from ai_examiner.services import document_ingestion
    from ai_examiner.services.storage import StorageService

    monkeypatch.setattr("ai_examiner.services.jobs.dispatch_job_by_id", lambda *args: None)
    project = client.post("/api/projects", json={"name": "Upload race"}).json()
    key = str(uuid4())
    url = f"/api/projects/{project['id']}/documents/async"
    data = {"request_id": key}
    files = {"file": ("paper.txt", b"Synthetic material", "text/plain")}
    first = client.post(url, data=data, files=files).json()
    original = document_ingestion.ingestion_receipt
    calls = []

    def miss_initial_read(*args):
        calls.append(1)
        return None if len(calls) == 1 else original(*args)

    monkeypatch.setattr(document_ingestion, "ingestion_receipt", miss_initial_read)
    second = client.post(url, data=data, files=files)
    assert second.status_code == 202, second.text
    assert second.json()["id"] == first["id"]
    with SessionLocal() as db:
        assert db.query(StoredObject).count() == 1
        source = db.scalar(select(StoredObject))
        assert StorageService(db, get_settings()).read_bytes(source) == b"Synthetic material"


def test_stale_retry_cannot_reset_running_worker(client, monkeypatch):
    from ai_examiner.services.job_control import reset_job_for_retry

    monkeypatch.setattr("ai_examiner.services.jobs.dispatch_job_by_id", lambda *args: None)
    project = client.post("/api/projects", json={"name": "Retry race"}).json()
    job_id = client.post(f"/api/projects/{project['id']}/documents/async", data={"request_id": str(uuid4())},
                         files={"file": ("paper.txt", b"Synthetic material", "text/plain")}).json()["id"]
    with SessionLocal() as db:
        job = db.get(BackgroundJob, job_id)
        job.status = "failed"
        db.commit()
        db.execute(update(BackgroundJob).where(BackgroundJob.id == job_id).values(
            status="running", lease_token="another-worker",
        ).execution_options(synchronize_session=False))
        db.commit()
        with pytest.raises(ValueError, match="Job changed"):
            reset_job_for_retry(db, job, reason="Stale retry request")
        db.refresh(job)
        assert job.status == "running" and job.lease_token == "another-worker"
