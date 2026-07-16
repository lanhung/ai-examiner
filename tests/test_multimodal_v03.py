from __future__ import annotations

import io

import fitz
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches


def _project(client, name="v0.3 multimodal"):
    response = client.post("/api/projects", json={"name": name})
    assert response.status_code == 201
    return response.json()["id"]


def _pdf_bytes() -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 80), "Multimodal Evidence Study", fontsize=18)
    page.insert_text(
        (72, 130),
        "Figure 1 compares the proposed examiner with a prompt-only baseline. "
        "The evidence is limited to text sessions and does not establish causal superiority.",
        fontsize=11,
    )
    page.draw_rect(fitz.Rect(90, 220, 500, 480), color=(0, 0, 0))
    page.insert_text((110, 260), "Accuracy: Proposed 0.86, Baseline 0.72", fontsize=13)
    return doc.tobytes()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8), Inches(0.8))
    title.text = "Defense overview"
    body = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8), Inches(2))
    body.text = "Contribution: evidence-grounded questioning\nLimitation: no speech evaluation yet"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(6), Inches(1.2)).table
    table.cell(0, 0).text = "Model"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Proposed"
    table.cell(1, 1).text = "0.86"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_heading("Supplementary evidence", level=1)
    document.add_paragraph("The supplementary experiment covers only one domain.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Setting"
    table.cell(0, 1).text = "Result"
    table.cell(1, 0).text = "OOD"
    table.cell(1, 1).text = "Not tested"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_pdf_evidence_visual_and_highlight(client):
    project_id = _project(client)
    uploaded = client.post(
        f"/api/projects/{project_id}/documents",
        files={"file": ("paper.pdf", _pdf_bytes(), "application/pdf")},
    )
    assert uploaded.status_code == 201, uploaded.text
    document = uploaded.json()
    assert document["document_kind"] == "pdf"
    assert document["page_count"] == 1
    assert document["evidence_count"] >= 2

    evidence = client.get(f"/api/documents/{document['id']}/evidence")
    assert evidence.status_code == 200
    assets = evidence.json()["assets"]
    page_asset = next(asset for asset in assets if asset["kind"] == "page")
    text_asset = next(asset for asset in assets if asset["kind"] != "page" and asset["bbox"])
    assert client.get(page_asset["file_url"]).status_code == 200
    assert client.get(f"/api/evidence/{text_asset['id']}/highlight").status_code == 200

    analyzed = client.post(
        f"/api/documents/{document['id']}/visual-analyses",
        json={"profile": "mock:heuristic-v2", "max_pages": 1, "asynchronous": False},
    )
    assert analyzed.status_code == 202, analyzed.text
    payload = analyzed.json()
    assert payload["count"] == 1
    assert payload["analyses"][0]["data"]["exam_questions"]


def test_pptx_docx_and_joint_analysis(client):
    project_id = _project(client, "joint package")
    pptx = client.post(
        f"/api/projects/{project_id}/documents",
        files={
            "file": (
                "defense.pptx",
                _pptx_bytes(),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
    )
    docx = client.post(
        f"/api/projects/{project_id}/documents",
        files={
            "file": (
                "supplement.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert pptx.status_code == 201, pptx.text
    assert docx.status_code == 201, docx.text
    assert pptx.json()["document_kind"] == "presentation"
    assert docx.json()["document_kind"] == "document"

    joined = client.post(
        f"/api/projects/{project_id}/joint-analyses",
        json={
            "document_ids": [pptx.json()["id"], docx.json()["id"]],
            "profile": "mock:heuristic-v2",
        },
    )
    assert joined.status_code == 201, joined.text
    assert joined.json()["data"]["high_risk_questions"]


def test_prompt_registry_dataset_freeze_diff_and_async_job(client):
    prompts = client.get("/api/prompts")
    assert prompts.status_code == 200
    assert "visual_evidence" in prompts.json()["active_manifest"]
    created_prompt = client.post(
        "/api/prompts",
        json={
            "name": "visual_evidence",
            "role": "multimodal_examiner",
            "content": "Review visible evidence, cite boundaries, and ask one focused defense question.",
            "activate": True,
        },
    )
    assert created_prompt.status_code == 201
    assert created_prompt.json()["status"] == "active"

    project_id = _project(client, "dataset lifecycle")
    material = ("The system separates planning, analysis, policy, and evidence scoring. " * 20).encode()
    document_id = client.post(
        f"/api/projects/{project_id}/documents",
        files={"file": ("paper.txt", material, "text/plain")},
    ).json()["id"]
    first = client.post(
        f"/api/projects/{project_id}/golden-datasets",
        json={"document_id": document_id, "question_count": 4},
    ).json()
    second = client.post(
        f"/api/projects/{project_id}/golden-datasets/async",
        json={"document_id": document_id, "question_count": 5},
    )
    assert second.status_code == 202, second.text
    job = client.get(f"/api/jobs/{second.json()['id']}").json()
    assert job["status"] == "completed"
    second_dataset_id = job["result"]["dataset_id"]

    frozen = client.patch(
        f"/api/golden-datasets/{first['id']}/status", json={"status": "frozen"}
    )
    assert frozen.status_code == 200
    assert frozen.json()["status"] == "frozen"
    diff = client.get(f"/api/golden-datasets/{first['id']}/diff/{second_dataset_id}")
    assert diff.status_code == 200
    assert "summary" in diff.json()
