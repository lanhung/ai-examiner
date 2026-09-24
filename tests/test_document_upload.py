import asyncio
import io
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from docx import Document as WordDocument

from ai_examiner import main
from ai_examiner.config import get_settings


def word_bytes():
    document = WordDocument()
    document.add_heading("Synthetic thesis defense", 0)
    document.add_paragraph("The experimental design uses a held-out test set.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Baseline"
    table.cell(1, 0).text = "Proposed method"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def project(client):
    response = client.post("/api/projects", json={"name": "Upload regression"})
    assert response.status_code == 201
    return response.json()["id"]


def test_word_upload_keeps_original_filename_and_can_be_rediscovered(client):
    pid = project(client)
    response = client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("wx-temp-file", word_bytes(), "application/octet-stream")},
        data={"original_filename": "答辩论文.docx"},
    )
    assert response.status_code == 201, response.text
    data = response.json()
    assert data["filename"] == "答辩论文.docx"
    assert data["document_kind"] == "document"
    assert data["evidence_count"] >= 3
    listed = client.get(f"/api/projects/{pid}/documents").json()
    assert listed[0]["id"] == data["id"]
    evidence = client.get(f"/api/documents/{data['id']}/evidence").json()
    preview = next(a for a in evidence["assets"] if a["has_file"])
    assert client.get(preview["file_url"]).status_code == 200
    assert list((get_settings().upload_dir / pid).iterdir()) == []


@pytest.mark.parametrize("name", ["legacy.doc", "../legacy.doc", "bad.exe"])
def test_unsupported_file_is_rejected_without_leftover_staging(client, name):
    pid = project(client)
    response = client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("wx-temp", b"invalid data")},
        data={"original_filename": name},
    )
    assert response.status_code == 400
    assert list((get_settings().upload_dir / pid).iterdir()) == []
    assert client.get(f"/api/projects/{pid}/documents").json() == []


def test_file_size_limit_and_filename_path_safety(client, monkeypatch):
    pid = project(client)
    monkeypatch.setattr(get_settings(), "max_upload_mb", 1)
    response = client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("too-big.txt", b"x" * (1024 * 1024 + 1))},
    )
    assert response.status_code == 400
    name = "../" + "a" * 200 + ".txt"
    response = client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("temp", b"Synthetic short evidence.")},
        data={"original_filename": name},
    )
    assert response.status_code == 201, response.text
    assert Path(response.json()["filename"]).name == response.json()["filename"]
    assert response.json()["filename"].endswith(".txt")


def test_slow_parser_does_not_block_health_requests(client, monkeypatch):
    pid = project(client)
    entered, release = threading.Event(), threading.Event()
    real_parse = main.parse_uploaded_document

    def slow_parse(*args, **kwargs):
        # A worker must have no running event loop. This catches async endpoint regressions.
        with pytest.raises(RuntimeError):
            asyncio.get_running_loop()
        entered.set()
        assert release.wait(10)
        return real_parse(*args, **kwargs)

    monkeypatch.setattr(main, "parse_uploaded_document", slow_parse)
    with ThreadPoolExecutor(max_workers=2) as pool:
        upload = pool.submit(
            client.post, f"/api/projects/{pid}/documents",
            files={"file": ("paper.docx", word_bytes())},
        )
        try:
            assert entered.wait(5)
            assert pool.submit(client.get, "/health").result(timeout=3).status_code == 200
        finally:
            release.set()
        assert upload.result(timeout=10).status_code == 201
